# -*- coding: utf-8 -*-
"""
지표 로그 검증 로직 (crawl_log_check_gui.py 에서 사용)

1) 로그정보 문자열 파싱 ({key=value, ...})
2) 로그정의서(검수 엑셀) 로드 → 로그별 정의 행 매칭
3) 검증
   - 정의 여부 (미정의 / 삭제된 로그)
   - 필드 규칙 (정의서에 값이 정의된 필드가 로그에 비어 있는지)
   - 수치 계산 (이전 ± 변동 = 최종)
   - 수치 연속성 (같은 플레이어·같은 재화/아이템에서 직전 최종값 = 이번 이전값)
4) 결과를 정의서 사본의 AOS/IOS 결과 칸에 기입
"""

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

CONFIG_PATH = Path(__file__).with_name("crawl_log_check_config.json")

# 재화 코드(rCurrency) → 정의서 재화명. 게임마다 다를 수 있어 설정 파일/GUI에서 수정 가능
DEFAULT_CONFIG = {
    # Start Browser 시 여는 어드민 로그 조회 페이지 (비어 있으면 실행 시 입력받아 저장)
    "admin_url": "",
    "currency_map": {
        "FREE_GOLD": "골드",
        "PAID_GOLD": "골드",
        "FREE_GEM": "스타젬",
        "PAID_GEM": "스타젬",
        "HEART": "하트",
    },
    # valueNoAfter = valueNoBefore + valueNo 로 계산 검증할 action reason
    # (정의서에서 valueNo=변동량, valueNoAfter=현재, valueNoBefore=직전 으로 정의된 것들)
    "action_value_reasons": [
        "levelUp",
        "ex_gain",
        "toploader_like",
        "toploader_report",
        "toploader_cheer",
        "toploader_cheer_cancle",
        "toploader_share",
    ],
    # 정의서에 값이 적혀 있어도 누락 검사에서 제외할 필드
    # itemName: 실제 로그에 필드 자체가 없음 / itemGrade: 카드류에만 값이 있음
    "ignore_missing_fields": ["itemName", "itemGrade"],
}

# 로그타입(크롤링 결과) → 정의서 시트 종류
LOGTYPE_KIND = {"item": "item", "cashitem": "resource", "resource": "resource", "action": "action"}
LOGTYPE_CLASS = {"cashitem": "CashItemLog", "resource": "ResourceLog"}

DEF_SHEETS = {"writeResourceLog": "resource", "writeItemLog": "item", "writeActionLog": "action"}

# 정의서에 값(설명)이 적혀 있으면 로그에도 값이 있어야 하는 필드
EXPECT_FIELDS = {
    "resource": ["subReason", "resourceAttr1", "resourceAttr2", "delta", "amount", "amountBefore"],
    "item": ["subReason", "itemAttr1", "itemAttr2", "quantity", "quantityBefore", "quantityAfter",
             "itemId", "itemType", "itemName", "itemGrade"],
    "action": ["subReason", "actionAttr1", "actionAttr2", "valueStr",
               "valueNo", "valueNoAfter", "valueNoBefore"],
}

RESULT_PASS = "Pass"
RESULT_FAIL = "Fail"
RESULT_WARN = "경고"
RESULT_UNDEF = "미정의"
RESULT_NA = "대상아님"


# ---------------------------------------------------------------------------
# 설정
# ---------------------------------------------------------------------------
def load_config(path=CONFIG_PATH) -> dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))  # deep copy
    path = Path(path)
    if path.exists():
        try:
            user = json.loads(path.read_text(encoding="utf-8"))
            for k, v in user.items():
                cfg[k] = v
        except Exception:
            pass  # 깨진 설정 파일이면 기본값 사용
    else:
        save_config(cfg, path)
    return cfg


def save_config(cfg: dict, path=CONFIG_PATH):
    Path(path).write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------------------------------------------------------------------
# 로그정보 파싱
# ---------------------------------------------------------------------------
_KEY_AT = re.compile(r"\s*[A-Za-z_][A-Za-z0-9_]*=")


def parse_loginfo(text: str) -> dict:
    """'{a=1, b={x=1, y=2}, c=[..], clientIp=1.1.1.1, 2.2.2.2}' → {'a': '1', 'b': '{x=1, y=2}', ...}

    - 최상위 레벨의 ', key=' 에서만 분리 (중괄호/대괄호 안의 쉼표는 무시)
    - 'clientIp=1.1.1.1, 2.2.2.2' 처럼 쉼표 뒤가 key= 형태가 아니면 값의 일부로 취급
    """
    if not text:
        return {}
    s = text.strip()
    if s.startswith("{") and s.endswith("}"):
        s = s[1:-1]

    parts, depth, start = [], 0, 0
    for i, ch in enumerate(s):
        if ch in "{[":
            depth += 1
        elif ch in "}]":
            depth = max(0, depth - 1)
        elif ch == "," and depth == 0 and _KEY_AT.match(s, i + 1):
            parts.append(s[start:i])
            start = i + 1
    parts.append(s[start:])

    fields = {}
    for p in parts:
        k, sep, v = p.strip().partition("=")
        if sep and k.strip():
            fields.setdefault(k.strip(), v.strip())
    return fields


def to_num(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return float(s)
        except ValueError:
            return None


def _fmt(n):
    return f"{n:,}" if isinstance(n, int) else str(n)


def _cell_str(v) -> str:
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------------------
# 로그정의서
# ---------------------------------------------------------------------------
class DefRow:
    def __init__(self, sheet, row, kind):
        self.sheet = sheet
        self.row = row
        self.kind = kind
        self.state = ""
        self.log_class = ""   # resource: CashItemLog / ResourceLog
        self.currency = ""    # resource: 골드/스타젬/하트 ...
        self.modType = ""
        self.reason = ""
        self.reasonNm = ""
        self.category = ""
        self.label = ""
        self.action = ""
        self.expected = {}    # 필드 → 정의서 설명

    @property
    def ref(self):
        return f"{self.sheet}!{self.row}"

    @property
    def deleted(self):
        return self.state == "삭제"

    def title(self):
        if self.kind == "resource":
            return f"{self.currency}/{self.modType}/{self.reason}"
        if self.kind == "item":
            return f"{self.modType}/{self.reason}"
        return f"{self.category}/{self.label}/{self.action}/{self.reason}"


class LogDefinition:
    def __init__(self, path):
        self.path = str(path)
        self.rows = []                 # [DefRow]
        self.sheet_cols = {}           # sheet → {header: col}
        self.sheet_header_row = {}     # sheet → header row
        self._load()

    # -------- 로드 --------
    @staticmethod
    def _find_header(ws):
        """'reason' 이 있는 행을 헤더로 보고, 바로 윗행의 제목(상태, AOS/IOS 등)도 함께 등록
        (writeActionLog 는 AOS/IOS 제목이 헤더 윗행에 있고, 헤더 행 같은 칸엔 'Not Test' 가 있음)"""
        for r in range(1, 21):
            vals = [_cell_str(c.value) for c in ws[r]]
            if "reason" in vals:
                above = [_cell_str(c.value) for c in ws[r - 1]] if r > 1 else []
                cols = {}
                for idx in range(len(vals)):
                    for name in (vals[idx], above[idx] if idx < len(above) else ""):
                        if name and name not in cols:
                            cols[name] = idx + 1
                return r, cols
        return None, {}

    def _load(self):
        wb = openpyxl.load_workbook(self.path, data_only=True)
        for sheet, kind in DEF_SHEETS.items():
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            hdr, cols = self._find_header(ws)
            if not hdr:
                continue
            self.sheet_cols[sheet] = cols
            self.sheet_header_row[sheet] = hdr

            # 병합 셀로 묶인 행은 reason 이 첫 행에만 있음 → 아래 행이 비어 있으면 물려받음
            inherit_from = {}
            for rng in ws.merged_cells.ranges:
                for r in range(rng.min_row + 1, rng.max_row + 1):
                    inherit_from.setdefault(r, rng.min_row)

            def get(r, name):
                c = cols.get(name)
                return _cell_str(ws.cell(r, c).value) if c else ""

            by_row = {}
            for r in range(hdr + 2, ws.max_row + 1):   # 헤더 다음 행은 설명 행
                reason = get(r, "reason")
                if not reason and r in inherit_from and inherit_from[r] in by_row:
                    reason = by_row[inherit_from[r]].reason
                if not reason:
                    continue

                d = DefRow(sheet, r, kind)
                d.state = get(r, "상태")
                d.reason = reason
                d.reasonNm = get(r, "reasonNm")
                d.modType = get(r, "modType")
                if kind == "resource":
                    d.log_class = get(r, "로그 분류")
                    d.currency = get(r, "rCurrency")
                if kind == "action":
                    d.category = get(r, "category")
                    d.label = get(r, "label")
                    d.action = get(r, "action")

                for f in EXPECT_FIELDS[kind]:
                    desc = get(r, f)
                    if f == "subReason" and not desc:
                        desc = get(r, "subReasonNm")
                    if desc:
                        d.expected[f] = desc
                self.rows.append(d)
                by_row[r] = d

    # -------- 조회 --------
    def reasons(self):
        return {d.reason for d in self.rows}

    def currencies(self):
        return sorted({d.currency for d in self.rows if d.currency})

    def result_cols(self, sheet):
        cols = self.sheet_cols.get(sheet, {})
        return cols.get("AOS"), cols.get("IOS")

    def match(self, logtype, fields, cfg):
        """→ (DefRow | None, 힌트 문자열)"""
        kind = LOGTYPE_KIND.get(logtype)
        reason = fields.get("reason", "")
        modType = fields.get("modType", "")
        cands = [d for d in self.rows if d.kind == kind and d.reason == reason]

        if kind == "resource":
            code = fields.get("rCurrency", "")
            label = cfg.get("currency_map", {}).get(code, "")
            same_mod = [d for d in cands if d.modType == modType]
            hit = [d for d in same_mod if d.currency == label]
            if not hit:
                if not label:
                    return None, f"재화 매핑 없음: rCurrency={code} (재화 매핑 설정 필요)"
                if same_mod:
                    have = ", ".join(sorted({d.currency for d in same_mod}))
                    return None, f"정의서에 '{label}/{modType}/{reason}' 없음 (같은 reason 정의 재화: {have})"
                if cands:
                    return None, f"정의서에 modType={modType} 없음 (정의: {', '.join(sorted({d.modType for d in cands}))})"
                return None, f"정의서에 reason={reason} 없음"
            cands = hit

        elif kind == "item":
            hit = [d for d in cands if d.modType == modType]
            if not hit:
                if cands:
                    return None, f"정의서에 modType={modType} 없음 (정의: {', '.join(sorted({d.modType for d in cands}))})"
                return None, f"정의서에 reason={reason} 없음"
            cands = hit

        elif kind == "action":
            if not cands:
                key = (fields.get("category", ""), fields.get("label", ""), fields.get("action", ""))
                similar = [d.reason for d in self.rows
                           if d.kind == "action" and (d.category, d.label, d.action) == key]
                hint = f" (같은 {key[0]}/{key[1]}/{key[2]} 정의 reason: {', '.join(similar)})" if similar else ""
                return None, f"정의서에 reason={reason} 없음{hint}"
            exact = [d for d in cands if d.label == fields.get("label") and d.action == fields.get("action")]
            cands = exact or cands

        if not cands:
            return None, "정의서에 없음"
        alive = [d for d in cands if not d.deleted]
        return (alive or cands)[0], ""


# ---------------------------------------------------------------------------
# 검증
# ---------------------------------------------------------------------------
def _check_numbers(row, cfg):
    """단건 수치 계산 검증 → [(level, msg)]"""
    f = row["fields"]
    t = row["로그타입"]
    issues = []
    if t == "item":
        q, b, a = to_num(f.get("quantity")), to_num(f.get("quantityBefore")), to_num(f.get("quantityAfter"))
        mt = f.get("modType")
        if None not in (q, b, a) and mt in ("add", "sub"):
            exp = b + q if mt == "add" else b - q
            if exp != a:
                op = "+" if mt == "add" else "-"
                issues.append(("fail", f"계산 불일치: quantityBefore {_fmt(b)} {op} quantity {_fmt(q)} = {_fmt(exp)} ≠ quantityAfter {_fmt(a)}"))
    elif t in ("cashitem", "resource"):
        d, a, b = to_num(f.get("delta")), to_num(f.get("amount")), to_num(f.get("amountBefore"))
        mt = f.get("modType")
        if None not in (d, a, b):
            if b + d != a:
                issues.append(("fail", f"계산 불일치: amountBefore {_fmt(b)} + delta {_fmt(d)} = {_fmt(b + d)} ≠ amount {_fmt(a)}"))
            if (mt == "add" and d < 0) or (mt == "sub" and d > 0):
                issues.append(("fail", f"modType({mt})와 delta 부호 불일치: delta={_fmt(d)}"))
    elif t == "action" and f.get("reason") in cfg.get("action_value_reasons", []):
        n, a, b = to_num(f.get("valueNo")), to_num(f.get("valueNoAfter")), to_num(f.get("valueNoBefore"))
        if None not in (n, a, b) and b + n != a:
            issues.append(("fail", f"계산 불일치: valueNoBefore {_fmt(b)} + valueNo {_fmt(n)} = {_fmt(b + n)} ≠ valueNoAfter {_fmt(a)}"))
    return issues


def _chain_key_and_values(row):
    f = row["fields"]
    t = row["로그타입"]
    pid = f.get("playerId", "")
    if t == "item":
        return (pid, "item", f.get("itemId", "")), to_num(f.get("quantityBefore")), to_num(f.get("quantityAfter"))
    if t in ("cashitem", "resource"):
        return (pid, "currency", f.get("rCurrency", "")), to_num(f.get("amountBefore")), to_num(f.get("amount"))
    return None, None, None


def _check_continuity(rows):
    """같은 플레이어·같은 재화/아이템 로그를 modTime 순으로 정렬해 직전 최종값 = 이번 이전값 확인"""
    groups = defaultdict(list)
    for r in rows:
        key, before, after = _chain_key_and_values(r)
        if key and key[2] and before is not None and after is not None:
            groups[key].append((r, before, after))

    for key, items in groups.items():
        items.sort(key=lambda x: (to_num(x[0]["fields"].get("modTime")) or 0))
        # 같은 modTime 끼리는 순서가 불확실 → 직전 최종값과 이어지는 것을 먼저 배치
        ordered = []
        i = 0
        while i < len(items):
            j = i
            t0 = items[i][0]["fields"].get("modTime")
            while j < len(items) and items[j][0]["fields"].get("modTime") == t0:
                j += 1
            tie = items[i:j]
            while tie:
                prev_after = ordered[-1][2] if ordered else None
                pick = next((x for x in tie if x[1] == prev_after), tie[0])
                ordered.append(pick)
                tie.remove(pick)
            i = j

        for idx, (r, before, after) in enumerate(ordered):
            if idx == 0:
                r["chain_note"] = "연속성 기준점(수집 범위의 첫 로그)"
                continue
            pr, _, p_after = ordered[idx - 1]
            if before != p_after:
                r["issues"].append((
                    "fail",
                    f"연속성 끊김: 직전 로그({pr['로그시간']} {pr['reason']}) 최종값 {_fmt(p_after)} ≠ 이번 이전값 {_fmt(before)}",
                ))


def verify_rows(rows, definition, cfg):
    """rows(크롤링 결과 dict 리스트)에 result/issues/def/detail 을 채운다"""
    ignore = set(cfg.get("ignore_missing_fields", []))

    for row in rows:
        row["issues"] = []
        row["def"] = None
        row["chain_note"] = ""
        t = row["로그타입"]
        f = row["fields"]
        kind = LOGTYPE_KIND.get(t)
        if not kind:
            continue

        if definition is not None:
            d, hint = definition.match(t, f, cfg)
            row["def"] = d
            if d is None:
                row["issues"].append(("undef", hint))
            else:
                if d.deleted:
                    row["issues"].append(("fail", f"정의서상 '삭제'된 로그가 발생 ({d.ref})"))
                if kind == "resource" and d.log_class and LOGTYPE_CLASS.get(t) != d.log_class:
                    row["issues"].append(("warn", f"로그 분류 불일치: 정의서 {d.log_class} / 실제 {t}"))
                if kind == "action":
                    for k, exp in (("category", d.category), ("label", d.label), ("action", d.action)):
                        if exp and f.get(k, "") != exp:
                            row["issues"].append(("fail", f"{k} 불일치: 정의서 {exp} / 실제 {f.get(k, '(없음)')}"))
                for fld, desc in d.expected.items():
                    if fld in ignore:
                        continue
                    if not f.get(fld, "").strip():
                        row["issues"].append(("fail", f"필드 누락: {fld} (정의: {desc})"))

        row["issues"].extend(_check_numbers(row, cfg))

    _check_continuity([r for r in rows if LOGTYPE_KIND.get(r["로그타입"])])

    for row in rows:
        if not LOGTYPE_KIND.get(row["로그타입"]):
            row["result"] = RESULT_NA
            row["detail"] = ""
            continue
        levels = {lv for lv, _ in row["issues"]}
        if "undef" in levels:
            row["result"] = RESULT_UNDEF
        elif "fail" in levels:
            row["result"] = RESULT_FAIL
        elif "warn" in levels:
            row["result"] = RESULT_WARN
        else:
            row["result"] = RESULT_PASS
        row["detail"] = " / ".join(msg for _, msg in row["issues"])


def os_key(row):
    """os 필드 기준, 없으면(서버 로그 등) market 으로 판단 → 'AOS' / 'IOS' / ''"""
    f = row["fields"]
    os_ = f.get("os", "").upper()
    if os_ == "ANDROID":
        return "AOS"
    if os_ == "IOS":
        return "IOS"
    market = f.get("market", "").lower()
    if market in ("googleplay", "google", "onestore", "galaxystore"):
        return "AOS"
    if market in ("apple", "appstore", "ios"):
        return "IOS"
    return ""


# ---------------------------------------------------------------------------
# 정의서 사본에 결과 기입
# ---------------------------------------------------------------------------
FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")
FILL_HEAD = PatternFill("solid", fgColor="DDEBF7")
LOG_SHEET_NAME = "자동검수_로그"


def write_results(src_path, dst_path, rows, definition):
    """정의서 사본(dst)에 결과 기입. 원본(src)은 수정하지 않음. → 요약 dict"""
    if Path(src_path).resolve() == Path(dst_path).resolve():
        raise ValueError("원본 정의서와 같은 파일에는 저장할 수 없습니다. 다른 이름으로 저장해주세요.")

    wb = openpyxl.load_workbook(src_path)   # 수식 유지 (data_only 아님)
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    agg = defaultdict(lambda: defaultdict(list))   # (sheet,row) → os → [rows]
    for r in rows:
        d = r.get("def")
        if d is not None and r.get("result") in (RESULT_PASS, RESULT_FAIL, RESULT_WARN):
            agg[(d.sheet, d.row)][os_key(r) or "OS미상"].append(r)

    detail_cols = {}
    for sheet in {s for s, _ in agg}:
        ws = wb[sheet]
        col = ws.max_column + 1
        hdr = definition.sheet_header_row[sheet]
        c = ws.cell(hdr, col, f"자동검수 상세 ({stamp})")
        c.fill = FILL_HEAD
        c.font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 60
        detail_cols[sheet] = col

    summary = {"rows": 0, "pass": 0, "fail": 0}
    for (sheet, rnum), per_os in sorted(agg.items()):
        ws = wb[sheet]
        aos_col, ios_col = definition.result_cols(sheet)
        lines = []
        for osk in ("AOS", "IOS", "OS미상"):
            lst = per_os.get(osk)
            if not lst:
                continue
            fails = [x for x in lst if x["result"] == RESULT_FAIL]
            res = RESULT_FAIL if fails else RESULT_PASS
            col = {"AOS": aos_col, "IOS": ios_col}.get(osk)
            if col:
                cell = ws.cell(rnum, col, res)
                cell.fill = FILL_FAIL if fails else FILL_PASS
                summary["fail" if fails else "pass"] += 1
            line = f"[{osk}] {res} (로그 {len(lst)}건" + (f", 실패 {len(fails)}건)" if fails else ")")
            msgs = []
            for x in fails:
                for lv, m in x["issues"]:
                    if lv == "fail" and m not in msgs:
                        msgs.append(m)
            if msgs:
                line += "\n  - " + "\n  - ".join(msgs[:5])
            lines.append(line)
        c = ws.cell(rnum, detail_cols[sheet], "\n".join(lines))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        summary["rows"] += 1

    # 전체 로그 + 검증 결과 시트
    if LOG_SHEET_NAME in wb.sheetnames:
        del wb[LOG_SHEET_NAME]
    ws = wb.create_sheet(LOG_SHEET_NAME, 1)
    headers = ["로그시간", "로그타입", "reason", "OS", "검증결과", "검증상세", "정의서 위치", "log_tx", "로그정보"]
    ws.append(headers)
    for c in ws[1]:
        c.fill = FILL_HEAD
        c.font = Font(bold=True)
    for r in rows:
        d = r.get("def")
        ws.append([
            r["로그시간"], r["로그타입"], r["reason"], os_key(r), r.get("result", ""),
            r.get("detail", ""), d.ref if d else "", r["fields"].get("log_tx", ""), r["로그정보"],
        ])
        res = r.get("result")
        if res in (RESULT_FAIL, RESULT_UNDEF):
            ws.cell(ws.max_row, 5).fill = FILL_FAIL
    for col, w in zip("ABCDEFGHI", (20, 11, 24, 6, 9, 70, 22, 16, 80)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(dst_path)
    return summary
